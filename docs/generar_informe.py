#!/usr/bin/env python3
"""Completa la plantilla institucional y genera el informe de la Actividad 8."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/home/borys/Downloads/Actividad_n_nombre_estudiante.docx")
OUTPUT = ROOT / "docs/Entrega_Actividad_8_Borys_Alexy_Panezo_Guerrero.docx"
EVIDENCIAS = ROOT / "docs/evidencias"

doc = Document(TEMPLATE)

reemplazos = {
    "4 créditos": "3 créditos",
    "Ing. Roly Steeven Cedeño Menéndez, Mtr.": "Ing. Jorge Arun Zambrano Cedeño, Mg.",
    "NOMBRES Y APELLIDOS: ": "NOMBRES Y APELLIDOS:  Borys Alexy Panezo Guerrero",
    "PARALELO:": "PARALELO:  A",
    "SEMESTRE:": "SEMESTRE:  Quinto",
    "ACTIVIDAD:": "ACTIVIDAD:  # 8 — Resolución de Ejercicios",
    "PERIODO SEPTIEMBRE 2025 - ENERO 2026": "PERIODO ABRIL - AGOSTO 2026",
}
for paragraph in doc.paragraphs:
    if paragraph.text in reemplazos:
        paragraph.text = reemplazos[paragraph.text]

# Retira únicamente el contenido de ejemplo posterior a la portada.
for start, paragraph in enumerate(list(doc.paragraphs)):
    if paragraph.text.startswith("Actividades a desarrollar"):
        for old in list(doc.paragraphs[start:]):
            old._element.getparent().remove(old._element)
        break

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.color.rgb = RGBColor(0, 130, 70)

def titulo(texto):
    p = doc.add_paragraph(texto, style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def parrafo(texto="", negrita=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(texto)
    r.bold = negrita
    return p

def codigo(texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(texto)
    r.font.name = "Liberation Mono"
    r.font.size = Pt(8)
    return p

def vineta(texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run("• ").bold = True
    p.add_run(texto)
    return p

doc.add_page_break()
titulo("Actividad # 8: Desarrollo de una API REST para Help Desk")
parrafo("Objetivo", True)
parrafo("Construir la arquitectura backend y la base de datos de un Sistema de Gestión de Incidentes, exponiendo una API RESTful segura que permita crear, consultar, actualizar y eliminar tickets.")

parrafo("Tecnologías utilizadas", True)
for texto in [
    "Java 21 y Spring Boot 3.5.16 para el servicio web.",
    "Spring Data JPA para persistencia y separación entre controlador, servicio y repositorio.",
    "PostgreSQL 17 como base de datos relacional.",
    "Flyway para crear y versionar la estructura de la tabla tickets.",
    "Spring Security con autenticación HTTP Basic y sesiones sin estado.",
    "JUnit, MockMvc y H2 para las pruebas automáticas.",
    "Docker para ejecutar la API y PostgreSQL en un entorno aislado.",
]:
    vineta(texto)

titulo("Arquitectura y estructura del proyecto")
parrafo("La solución utiliza una arquitectura por capas. El controlador recibe las peticiones HTTP; el servicio contiene la lógica de negocio y las transacciones; el repositorio gestiona la persistencia; y PostgreSQL conserva los datos. Las credenciales se reciben mediante variables de entorno y no forman parte del código fuente.")

tabla = doc.add_table(rows=1, cols=2)
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
tabla.style = "Table Grid"
tabla.rows[0].cells[0].text = "Componente"
tabla.rows[0].cells[1].text = "Responsabilidad"
for componente, responsabilidad in [
    ("TicketController", "Expone los cinco endpoints REST."),
    ("TicketService", "Aplica reglas, transacciones y tratamiento de datos."),
    ("TicketRepository", "Ejecuta las operaciones de persistencia JPA."),
    ("PostgreSQL", "Almacena tickets e índices de consulta."),
    ("ApiExceptionHandler", "Devuelve errores JSON uniformes."),
    ("SecurityConfig", "Protege todas las rutas mediante HTTP Basic."),
]:
    cells = tabla.add_row().cells
    cells[0].text, cells[1].text = componente, responsabilidad

titulo("Diseño de la base de datos")
parrafo("La migración V1 crea la tabla tickets con los atributos mínimos solicitados: id, título, descripción, categoría, prioridad y estado. También registra fechas de creación y actualización e índices para estado y prioridad.")
codigo("""CREATE TABLE tickets (
    id BIGSERIAL PRIMARY KEY,
    titulo VARCHAR(150) NOT NULL,
    descripcion VARCHAR(2000) NOT NULL,
    categoria VARCHAR(20) NOT NULL,
    prioridad VARCHAR(20) NOT NULL,
    estado VARCHAR(20) NOT NULL DEFAULT 'ABIERTO',
    creado_en TIMESTAMP WITH TIME ZONE NOT NULL,
    actualizado_en TIMESTAMP WITH TIME ZONE NOT NULL
);""")
parrafo("Valores admitidos: categoría RED, HARDWARE o SOFTWARE; prioridad ALTA, MEDIA o BAJA; estado ABIERTO, EN_PROGRESO o CERRADO.")

doc.add_page_break()
titulo("Implementación de la API REST")
tabla = doc.add_table(rows=1, cols=4)
tabla.style = "Table Grid"
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, texto in enumerate(["Método", "Ruta", "Resultado", "Código"]):
    tabla.rows[0].cells[i].text = texto
for fila in [
    ("GET", "/tickets", "Lista todos los tickets", "200"),
    ("GET", "/tickets/{id}", "Busca un ticket", "200/404"),
    ("POST", "/tickets", "Registra un ticket", "201/400"),
    ("PUT", "/tickets/{id}", "Actualiza un ticket", "200/400/404"),
    ("DELETE", "/tickets/{id}", "Elimina un ticket", "204/404"),
]:
    cells = tabla.add_row().cells
    for i, valor in enumerate(fila): cells[i].text = valor

parrafo("Código principal del controlador", True)
controller = (ROOT / "src/main/java/ec/edu/utm/helpdesk/ticket/TicketController.java").read_text(encoding="utf-8")
codigo(controller)

doc.add_page_break()
titulo("Pruebas y resultados")
parrafo("Las pruebas se ejecutaron contra PostgreSQL 17 en contenedores aislados. La API realizó correctamente el ciclo CRUD completo. Adicionalmente se ejecutaron pruebas automáticas con el siguiente resultado: 3 pruebas, 0 fallos, 0 errores y 0 omitidas.")
for nombre, leyenda in [
    ("01_crear.png", "Figura 1. Creación de un ticket — HTTP 201 Created."),
    ("02_listar.png", "Figura 2. Listado de tickets — HTTP 200 OK."),
    ("03_buscar.png", "Figura 3. Consulta de un ticket por identificador — HTTP 200 OK."),
    ("04_actualizar.png", "Figura 4. Actualización del estado del ticket — HTTP 200 OK."),
    ("05_eliminar.png", "Figura 5. Eliminación del ticket — HTTP 204 No Content."),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(EVIDENCIAS / nombre), width=Inches(6.3))
    cap = doc.add_paragraph(leyenda)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    if nombre != "05_eliminar.png": doc.add_page_break()

titulo("Conclusiones")
parrafo("La API desarrollada cumple los requisitos funcionales de la Actividad 8: utiliza PostgreSQL, implementa la estructura mínima de tickets, expone las cinco operaciones CRUD y devuelve información en JSON. La validación de entradas, los errores estandarizados, las migraciones y la autenticación mejoran la seguridad y mantenibilidad del servicio.")
parrafo("El proyecto puede ejecutarse localmente mediante Docker y contiene una colección Postman, pruebas automáticas y documentación de uso. El código se organizó en la rama feature/backend-api para integrarlo posteriormente en develop, siguiendo el flujo de control de versiones solicitado.")

titulo("Repositorio")
parrafo("Código fuente: https://github.com/JBorys687/helpdesk-datacenter")

titulo("Bibliografía")
for referencia in [
    "Spring. (2026). Spring Boot Reference Documentation. https://docs.spring.io/spring-boot/",
    "Spring. (2026). Spring Data JPA Reference Documentation. https://docs.spring.io/spring-data/jpa/reference/",
    "PostgreSQL Global Development Group. (2026). PostgreSQL Documentation. https://www.postgresql.org/docs/",
    "Fielding, R., Nottingham, M., & Reschke, J. (2022). HTTP Semantics (RFC 9110). https://www.rfc-editor.org/rfc/rfc9110",
]:
    vineta(referencia)

# Evita que Word solicite actualizar campos externos al abrir el documento.
settings = doc.settings.element
for node in settings.xpath("./w:updateFields"):
    settings.remove(node)

doc.save(OUTPUT)
print(OUTPUT)

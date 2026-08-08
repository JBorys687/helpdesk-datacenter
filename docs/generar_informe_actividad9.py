#!/usr/bin/env python3
"""Completa la plantilla institucional y genera el informe de la Actividad 9.

Antes de ejecutar, completar las constantes marcadas como PENDIENTE con los
datos reales tras el despliegue (URLs públicas, enlace del video, y colocar
las capturas de pantalla reales en docs/evidencias9/).
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/home/borys/Downloads/Actividad_n_nombre_estudiante.docx")
OUTPUT = ROOT / "docs/Entrega_Actividad_9_Borys_Alexy_Panezo_Guerrero.docx"
EVIDENCIAS = ROOT / "docs/evidencias9"

# ── Datos que dependen del despliegue (completar antes de generar la versión final) ──
URL_BACKEND = "https://helpdesk-datacenter-api.onrender.com"
URL_FRONTEND = "https://helpdesk-datacenter-blush.vercel.app"
URL_VIDEO = "PENDIENTE: enlace del video de sustentación (Drive/YouTube/Loom, sin restricción de acceso) — REQUISITO OBLIGATORIO, ver docs/GUION_VIDEO_SUSTENTACION.md"

doc = Document(TEMPLATE)

reemplazos = {
    "4 créditos": "3 créditos",
    "Ing. Roly Steeven Cedeño Menéndez, Mtr.": "Ing. Jorge Arun Zambrano Cedeño, Mg.",
    "NOMBRES Y APELLIDOS: ": "NOMBRES Y APELLIDOS:  Borys Alexy Panezo Guerrero",
    "PARALELO:": "PARALELO:  A",
    "SEMESTRE:": "SEMESTRE:  Quinto",
    "ACTIVIDAD:": "ACTIVIDAD:  # 9 — Proyecto Práctico: Frontend e Integración Full Stack",
    "PERIODO SEPTIEMBRE 2025 - ENERO 2026": "PERIODO ABRIL - AGOSTO 2026",
}
for paragraph in doc.paragraphs:
    if paragraph.text in reemplazos:
        paragraph.text = reemplazos[paragraph.text]

# Retira únicamente el contenido de ejemplo posterior a la portada.
for start, paragraph in enumerate(list(doc.paragraphs)):
    if paragraph.text.startswith("Actividades a desarrollar"):
        for old in list(doc.paragraphs[start:]):
            old._element.getparent().remove(old._element)
        break

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.color.rgb = RGBColor(0, 130, 70)

def titulo(texto):
    p = doc.add_paragraph(texto, style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def parrafo(texto="", negrita=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(texto)
    r.bold = negrita
    return p

def codigo(texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(texto)
    r.font.name = "Liberation Mono"
    r.font.size = Pt(8)
    return p

def vineta(texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run("• ").bold = True
    p.add_run(texto)
    return p

def parrafo_enlace(etiqueta, url):
    p = doc.add_paragraph()
    p.add_run(etiqueta + ": ")
    if url.startswith("PENDIENTE"):
        r = p.add_run(url)
        r.italic = True
        r.font.color.rgb = RGBColor(192, 57, 43)
        return p
    relacion = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), relacion)
    run = OxmlElement("w:r")
    propiedades = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    subrayado = OxmlElement("w:u")
    subrayado.set(qn("w:val"), "single")
    propiedades.append(color)
    propiedades.append(subrayado)
    texto = OxmlElement("w:t")
    texto.text = url
    run.append(propiedades)
    run.append(texto)
    enlace.append(run)
    p._p.append(enlace)
    return p

doc.add_page_break()
titulo("Actividad # 9: Desarrollo del Frontend e Integración Full Stack")
parrafo("Objetivo", True)
parrafo("Construir una Single Page Application (SPA) interactiva con Angular, aplicando modularización y consumiendo la API REST desarrollada en la Actividad 8, para completar un sistema web Full Stack desplegado en producción.")

parrafo("Tecnologías utilizadas", True)
for texto in [
    "Angular 22 (standalone components, signals, control flow @if/@for) para el frontend.",
    "TypeScript, RxJS y Reactive Forms para la lógica de la interfaz.",
    "HttpClient con un interceptor de autenticación HTTP Basic.",
    "Java 21 / Spring Boot / PostgreSQL / Flyway para el backend (Actividad 8).",
    "Docker para empaquetar el backend en producción.",
    "Render para desplegar el backend y la base de datos.",
    "Vercel para publicar el frontend como sitio estático.",
]:
    vineta(texto)

titulo("Arquitectura del sistema")
parrafo("El sistema sigue una arquitectura cliente/servidor desacoplada: el frontend Angular corre íntegramente en el navegador y se comunica por HTTPS con la API REST del backend, la cual persiste la información en una base de datos PostgreSQL gestionada de forma remota. Cada componente se despliega y escala de forma independiente.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(EVIDENCIAS / "diagrama_arquitectura.png"), width=Inches(6.3))
cap = doc.add_paragraph("Figura 1. Diagrama de arquitectura cliente/servidor.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True

titulo("Componentes del frontend")
tabla = doc.add_table(rows=1, cols=2)
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
tabla.style = "Table Grid"
tabla.rows[0].cells[0].text = "Componente"
tabla.rows[0].cells[1].text = "Responsabilidad"
for componente, responsabilidad in [
    ("Navegacion", "Barra de navegación fija, menú responsivo y cierre de sesión."),
    ("Login", "Captura las credenciales HTTP Basic y las guarda en sessionStorage."),
    ("Dashboard", "KPIs (total, por estado), gráficos de barra por categoría/prioridad y últimos tickets."),
    ("RegistroIncidentes", "Formulario reactivo con validaciones para crear tickets."),
    ("ListadoTickets", "Tabla con filtros, edición de estado y eliminación (CRUD completo)."),
    ("TicketService / Auth / authInterceptor / authGuard", "Capa de servicios: consumo de la API, sesión y protección de rutas."),
]:
    cells = tabla.add_row().cells
    cells[0].text, cells[1].text = componente, responsabilidad

doc.add_page_break()
titulo("Integración y lógica de negocio")
parrafo("El servicio TicketService centraliza el consumo de los cinco endpoints REST mediante peticiones HTTP asíncronas (HttpClient + RxJS). Un interceptor adjunta la cabecera Authorization (HTTP Basic) únicamente a las peticiones dirigidas a la API configurada, y un guard de rutas impide el acceso a las vistas internas sin sesión iniciada.")
parrafo("Prevención de XSS", True)
parrafo("Angular escapa automáticamente cualquier valor mostrado mediante interpolación ({{ }}); el proyecto no utiliza [innerHTML] en ningún componente, por lo que el contenido ingresado por el usuario (título, descripción del incidente) nunca se interpreta como HTML o JavaScript. Adicionalmente, los campos de texto se recortan (trim) y se limitan en longitud (Validators.maxLength) antes de enviarse al backend, y las credenciales solo se almacenan en sessionStorage del navegador (nunca en localStorage ni en el código fuente).")

parrafo("Código principal del servicio de tickets (frontend)", True)
ticket_service = (ROOT / "frontend/src/app/services/ticket.ts").read_text(encoding="utf-8")
codigo(ticket_service)

doc.add_page_break()
titulo("Despliegue en la nube")
tabla = doc.add_table(rows=1, cols=3)
tabla.style = "Table Grid"
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, texto in enumerate(["Componente", "Servicio", "URL pública"]):
    tabla.rows[0].cells[i].text = texto
for fila in [
    ("Base de datos (PostgreSQL)", "Render", "gestionada internamente"),
    ("Backend (API REST)", "Render (Docker)", URL_BACKEND),
    ("Frontend (SPA)", "Vercel", URL_FRONTEND),
]:
    cells = tabla.add_row().cells
    for i, valor in enumerate(fila): cells[i].text = valor
parrafo("El despliegue se automatizó con un Blueprint de Render (render.yaml), que aprovisiona la base de datos y el servicio web del backend a partir del Dockerfile del repositorio, y con Vercel para el sitio estático generado por el build de producción de Angular (vercel.json). Ambos servicios se conectan mediante variables de entorno: CORS_ALLOWED_ORIGINS en el backend habilita las peticiones desde el dominio del frontend, y environment.prod.ts en el frontend apunta a la URL pública del backend.")

doc.add_page_break()
titulo("Sustentación en video")
parrafo_enlace("Enlace al video de sustentación", URL_VIDEO)
parrafo("El video muestra la arquitectura del sistema, el código fuente principal, el proceso de despliegue realizado en Render y Vercel, y una demostración en vivo del sistema funcionando desde la URL pública (creación, actualización y eliminación de tickets).")

doc.add_page_break()
titulo("Capturas del sistema en producción")
parrafo("Capturas tomadas directamente desde las URLs públicas (no localhost), con un ticket real creado, listado y persistido en la base de datos remota.")
for nombre, leyenda in [
    ("dashboard_produccion.png", "Figura 2. Dashboard en producción (https://helpdesk-datacenter-blush.vercel.app/dashboard), con datos reales servidos por la API."),
    ("listado_tickets_produccion.png", "Figura 3. Listado de Tickets en producción, mostrando el incidente creado en vivo."),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(EVIDENCIAS / nombre), width=Inches(6.3))
    cap = doc.add_paragraph(leyenda)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    doc.add_page_break()

titulo("Conclusiones técnicas")
parrafo("La integración full-stack evidenció la importancia de separar responsabilidades entre frontend y backend desplegados en dominios distintos: fue necesario habilitar CORS explícitamente en el backend para que el navegador permitiera las peticiones del frontend en producción. El uso de un Blueprint (render.yaml) y de vercel.json permitió reproducir el entorno de despliegue de forma declarativa y documentada, cumpliendo el requisito de que el sistema sea 100% funcional en la web sin depender de localhost.")
parrafo("Angular, mediante su escape automático de interpolaciones y el uso de formularios reactivos con validación, redujo significativamente el riesgo de vulnerabilidades XSS sin necesidad de librerías adicionales de sanitización.")

titulo("Repositorio")
parrafo_enlace("Repositorio principal", "https://github.com/JBorys687/helpdesk-datacenter")
parrafo_enlace("Rama de desarrollo integrada (unifica backend y frontend)", "https://github.com/JBorys687/helpdesk-datacenter/tree/develop")
parrafo_enlace("Rama del frontend (Actividad 9)", "https://github.com/JBorys687/helpdesk-datacenter/tree/feature/frontend-app")
parrafo_enlace("URL pública del sistema (frontend)", URL_FRONTEND)
parrafo_enlace("URL pública de la API (backend)", URL_BACKEND)

titulo("Bibliografía")
for referencia in [
    "Angular Team. (2026). Angular Documentation. https://angular.dev/",
    "Spring. (2026). Spring Boot Reference Documentation. https://docs.spring.io/spring-boot/",
    "Render. (2026). Render Documentation — Blueprints. https://render.com/docs/blueprint-spec",
    "Vercel. (2026). Vercel Documentation. https://vercel.com/docs",
    "OWASP Foundation. (2026). Cross Site Scripting (XSS) Prevention Cheat Sheet. https://cheatsheetseries.owasp.org/",
]:
    vineta(referencia)

# Evita que Word solicite actualizar campos externos al abrir el documento.
settings = doc.settings.element
for node in settings.xpath("./w:updateFields"):
    settings.remove(node)

doc.save(OUTPUT)
print(OUTPUT)
